from datetime import date
from html import unescape
from io import BytesIO

from openpyxl import load_workbook

from odoo.tests import tagged

from odoo.addons.account.tests.common import AccountTestInvoicingCommon

REPORT_MODEL = "report.gc_odoo_account_report_custom.trial_balance_monthly"
HTML_REPORT_NAME = "gc_odoo_account_report_custom.trial_balance_monthly"
XLSX_REPORT_NAME = "gc_odoo_account_report_custom.trial_balance_monthly_xlsx"
AMOUNT_FIELDS = ("debit", "credit", "balance")


@tagged("post_install", "-at_install")
class TestTrialBalanceMonthly(AccountTestInvoicingCommon):
    @classmethod
    def setUpClass(cls, chart_template_ref=None):
        super().setUpClass(chart_template_ref=chart_template_ref)
        cls.report_model = cls.env[REPORT_MODEL]
        cls.group = cls.env["account.group"].create(
            {"code_prefix_start": "TBM", "name": "Grupo TBM"}
        )
        cls.account_asset = cls.env["account.account"].create(
            {"code": "TBM100", "name": "Activo TBM", "account_type": "asset_current"}
        )
        cls.account_income = cls.env["account.account"].create(
            {"code": "TBM200", "name": "Ingreso TBM", "account_type": "income_other"}
        )
        (cls.account_asset | cls.account_income).group_id = cls.group
        # Ejercicio anterior: saldo inicial de la cuenta patrimonial
        cls._add_move("2015-12-15", 1000)
        cls._add_move("2016-01-10", 100)
        cls._add_move("2016-03-05", 50)
        cls._add_move("2016-03-20", -20)
        cls._add_move("2016-02-10", 7, post=False)

    @classmethod
    def _add_move(cls, move_date, amount, post=True):
        """Asiento que debita `amount` en el activo contra el ingreso.

        Un importe negativo invierte el asiento.
        """
        debit, credit = (amount, 0.0) if amount > 0 else (0.0, -amount)
        move = cls.env["account.move"].create(
            {
                "journal_id": cls.company_data["default_journal_misc"].id,
                "date": move_date,
                "line_ids": [
                    (
                        0,
                        0,
                        {
                            "account_id": cls.account_asset.id,
                            "debit": debit,
                            "credit": credit,
                        },
                    ),
                    (
                        0,
                        0,
                        {
                            "account_id": cls.account_income.id,
                            "debit": credit,
                            "credit": debit,
                        },
                    ),
                ],
            }
        )
        if post:
            move.action_post()
        return move

    def _create_wizard(self, **vals):
        values = {
            "company_id": self.company_data["company"].id,
            "date_from": "2016-01-01",
            "date_to": "2016-03-31",
            "account_ids": [(6, 0, (self.account_asset | self.account_income).ids)],
        }
        values.update(vals)
        return self.env["trial.balance.report.wizard"].create(values)

    def _get_report_values(self, wizard):
        data = wizard._prepare_monthly_report_data()
        return self.report_model._get_report_values(wizard, data)

    def _get_line(self, res, line_type, record):
        return next(
            line
            for line in res["trial_balance"]
            if line["type"] == line_type and line["id"] == record.id
        )

    def _get_monthly_amounts(self, line, periods):
        return [
            tuple(
                line[self.report_model._get_monthly_field_name(index, field)]
                for field in AMOUNT_FIELDS
            )
            for index in range(len(periods))
        ]

    def test_monthly_periods_full_year(self):
        periods = self.report_model._get_monthly_periods("2016-01-01", "2016-12-31")
        self.assertEqual(len(periods), 12)
        self.assertEqual(periods[0]["date_from"], date(2016, 1, 1))
        self.assertEqual(periods[0]["date_to"], date(2016, 1, 31))
        self.assertEqual(periods[1]["date_to"], date(2016, 2, 29))
        self.assertEqual(periods[11]["date_from"], date(2016, 12, 1))
        self.assertEqual(periods[11]["date_to"], date(2016, 12, 31))

    def test_monthly_periods_are_clipped_to_date_range(self):
        periods = self.report_model._get_monthly_periods(
            date(2015, 12, 15), date(2016, 2, 10)
        )
        self.assertEqual(
            [(period["date_from"], period["date_to"]) for period in periods],
            [
                (date(2015, 12, 15), date(2015, 12, 31)),
                (date(2016, 1, 1), date(2016, 1, 31)),
                (date(2016, 2, 1), date(2016, 2, 10)),
            ],
        )
        self.assertEqual(len({period["name"] for period in periods}), 3)

    def test_account_amounts_are_split_by_month(self):
        res = self._get_report_values(self._create_wizard())
        self.assertEqual(len(res["periods"]), 3)
        line = self._get_line(res, "account_type", self.account_asset)
        self.assertEqual(line["initial_balance"], 1000)
        self.assertEqual(
            self._get_monthly_amounts(line, res["periods"]),
            [(100, 0, 100), (0, 0, 0), (50, 20, 30)],
        )
        self.assertEqual(line["debit"], 150)
        self.assertEqual(line["credit"], 20)
        self.assertEqual(line["ending_balance"], 1130)

    def test_draft_moves_are_included_when_all_entries_are_targeted(self):
        res = self._get_report_values(self._create_wizard(target_move="all"))
        line = self._get_line(res, "account_type", self.account_asset)
        self.assertEqual(
            self._get_monthly_amounts(line, res["periods"]),
            [(100, 0, 100), (7, 0, 7), (50, 20, 30)],
        )

    def test_monthly_amounts_add_up_to_period_totals(self):
        res = self._get_report_values(self._create_wizard(show_hierarchy=True))
        self.assertTrue(res["trial_balance"])
        for line in res["trial_balance"]:
            monthly = self._get_monthly_amounts(line, res["periods"])
            for position, field in enumerate(AMOUNT_FIELDS):
                self.assertAlmostEqual(
                    sum(amounts[position] for amounts in monthly),
                    line[field],
                    msg="%s / %s" % (line["name"], field),
                )

    def test_group_lines_sum_their_accounts_by_month(self):
        res = self._get_report_values(self._create_wizard(show_hierarchy=True))
        line = self._get_line(res, "group_type", self.group)
        self.assertEqual(
            self._get_monthly_amounts(line, res["periods"]),
            [(100, 100, 0), (0, 0, 0), (70, 70, 0)],
        )

    def test_wizard_exports_monthly_xlsx_without_unsupported_options(self):
        wizard = self._create_wizard(show_partner_details=True, foreign_currency=True)
        action = wizard.button_export_monthly_xlsx()
        self.assertEqual(action["report_name"], XLSX_REPORT_NAME)
        self.assertEqual(action["report_type"], "xlsx")
        self.assertFalse(action["data"]["show_partner_details"])
        self.assertFalse(action["data"]["foreign_currency"])
        self.assertFalse(action["data"]["grouped_by"])

    def test_xlsx_has_three_columns_per_month(self):
        wizard = self._create_wizard()
        data = wizard._prepare_monthly_report_data()
        content, file_type = self.env["ir.actions.report"]._render_xlsx(
            XLSX_REPORT_NAME, wizard.ids, data
        )
        self.assertEqual(file_type, "xlsx")
        rows = list(load_workbook(BytesIO(content)).active.iter_rows(values_only=True))
        periods = self.report_model._get_monthly_periods(
            wizard.date_from, wizard.date_to
        )
        header_index = next(
            index for index, row in enumerate(rows) if row[1] == "Cuenta"
        )
        month_row = rows[header_index - 1]
        for index, period in enumerate(periods):
            self.assertEqual(month_row[3 + index * 3], period["name"])
        self.assertEqual(
            rows[header_index],
            ("Código", "Cuenta", "Saldo inicial")
            + ("Debe", "Haber", "Saldo") * 4
            + ("Saldo final",),
        )
        asset_row = next(row for row in rows if row[0] == "TBM100")
        self.assertEqual(
            asset_row,
            ("TBM100", "Activo TBM", 1000)
            + (100, 0, 100, 0, 0, 0, 50, 20, 30)
            + (150, 20, 130, 1130),
        )

    def test_hierarchy_level_limit_filters_lines(self):
        wizard = self._create_wizard(
            show_hierarchy=True, limit_hierarchy_level=True, show_hierarchy_level=1
        )
        res = self._get_report_values(wizard)
        self.assertEqual(
            [(line["type"], line["id"]) for line in res["trial_balance"]],
            [("group_type", self.group.id)],
        )

    def test_wizard_opens_monthly_html_view(self):
        action = self._create_wizard().button_export_monthly_html()
        self.assertEqual(action["report_name"], HTML_REPORT_NAME)
        self.assertEqual(action["report_type"], "qweb-html")
        self.assertFalse(action["data"]["show_partner_details"])

    def test_html_view_shows_amounts_by_month(self):
        wizard = self._create_wizard()
        data = wizard._prepare_monthly_report_data()
        html = unescape(
            self.env["ir.actions.report"]
            ._render_qweb_html(HTML_REPORT_NAME, wizard.ids, data)[0]
            .decode()
        )
        for period in self.report_model._get_monthly_periods(
            wizard.date_from, wizard.date_to
        ):
            self.assertIn(period["name"], html)
        self.assertIn("TBM100", html)
        self.assertIn("1,130.00", html)
        # Los importes del mes abren sus apuntes
        self.assertIn("('date', '>=', '2016-03-01')", html)
        self.assertIn("('date', '<=', '2016-03-31')", html)

    def test_html_group_amounts_open_the_group_entries(self):
        wizard = self._create_wizard(show_hierarchy=True)
        data = wizard._prepare_monthly_report_data()
        html = unescape(
            self.env["ir.actions.report"]
            ._render_qweb_html(HTML_REPORT_NAME, wizard.ids, data)[0]
            .decode()
        )
        self.assertIn("('account_id.group_id', 'child_of', %d)" % self.group.id, html)
        lines = self.env["account.move.line"].search(
            [
                ("account_id.group_id", "child_of", self.group.id),
                ("parent_state", "=", "posted"),
                ("date", ">=", "2016-03-01"),
                ("date", "<=", "2016-03-31"),
            ]
        )
        self.assertEqual(sum(lines.mapped("debit")), 70)
